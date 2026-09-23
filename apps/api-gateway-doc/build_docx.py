"""Build the API Gateway Enterprise Architecture Standard (.docx).

    ../.venv/bin/python build_docx.py             # build the working copy
    ../.venv/bin/python build_docx.py --archive   # ...and archive it as versions/v<version>/
    ../.venv/bin/python build_docx.py --lang he   # Hebrew (right-to-left) edition, from content_he.py

Inputs : content.py (the document as data), content_he.py (its Hebrew translation), figures.py (diagrams)
Outputs: API_Gateway_Enterprise_Architecture_Standard.docx, CHANGES.md
         API_Gateway_Enterprise_Architecture_Standard_he.docx (Hebrew; no CHANGES.md; archived with --archive from 1.4)
Options: --out PATH   write the .docx there instead (used for regression checks)
"""
import os
import re
import shutil
import sys
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT,
                            WD_TAB_LEADER)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import content as C
import figures

C_EN = C        # the English content stays available as the reference structure for translations

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DOCX = os.path.join(HERE, "API_Gateway_Enterprise_Architecture_Standard.docx")
OUT_CHANGES = os.path.join(HERE, "CHANGES.md")
README = os.path.join(HERE, "README.md")

FONT, MONO = "Calibri", "Consolas"
NAVY, BLUE, GRAY = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x2F, 0x54, 0x96), RGBColor(0x59, 0x59, 0x59)
NAVY_HEX, BLUE_HEX, LIGHT_HEX, GRID_HEX = "1F3864", "2F5496", "EAF0F9", "BFBFBF"
TEXT_W_CM = 16.0
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
             "rPr", "sectPr", "pPrChange"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd",
               "tblLayout", "tblCellMar", "tblLook"]
SECTPR_ORDER = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
                "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
                "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings"]

errors = []

LANG, RTL = "en", False
HEBREW = re.compile(r"[\u0590-\u05FF]")

# Wording that is not in content.py: labels, headings and fixed phrases, per language.
STR = {
    "en": dict(
        font="Calibri", chapter="Chapter", appendix="Appendix", section="Section", figure="Figure",
        table="Table", contents="Contents", list_figures="List of Figures", list_tables="List of Tables",
        doc_control="Document Control", revision_history="Revision History", approval="Approval",
        note="Note. ", version="Version", status="Status", date="Date", doc_id="Document ID",
        classification="Classification", doc_title="Document title", owner="Document owner",
        author="Author", effective="Effective date", review_cycle="Review cycle",
        next_review="Next review date", rev_headers=["Version", "Date", "Author", "Description"],
        appr_headers=["Role", "Name", "Signature", "Date"], role="[[Role]]", name="[[Name]]",
        footer_version="Version {v} — {status}", page="Page", of=" of ",
        keywords="API Gateway; architecture; operating model; standard",
        comments="Draft. Converted from README.md.", appendix_letters={}),
    "he": dict(
        font="Arial", chapter="פרק", appendix="נספח", section="סעיף", figure="איור", table="טבלה",
        contents="תוכן עניינים", list_figures="רשימת איורים", list_tables="רשימת טבלאות",
        doc_control="בקרת מסמך", revision_history="היסטוריית גרסאות", approval="אישור",
        note="הערה. ", version="גרסה", status="סטטוס", date="תאריך", doc_id="מזהה מסמך",
        classification="סיווג", doc_title="שם המסמך", owner="בעל המסמך", author="מחבר",
        effective="תאריך כניסה לתוקף", review_cycle="מחזור סקירה", next_review="תאריך הסקירה הבאה",
        rev_headers=["גרסה", "תאריך", "מחבר", "תיאור"], appr_headers=["תפקיד", "שם", "חתימה", "תאריך"],
        role="[[תפקיד]]", name="[[שם]]", footer_version="גרסה {v} — {status}", page="עמוד", of=" מתוך ",
        keywords="API Gateway; ארכיטקטורה; מודל תפעולי; תקן",
        comments="טיוטה. תרגום לעברית של גרסה 1.4.",
        appendix_letters={"A": "א׳", "B": "ב׳", "C": "ג׳"}),
}
T = STR["en"]


def select_language(lang):
    """Switch the build to another language edition (module-level state, set once per run)."""
    global LANG, RTL, FONT, T, C, OUT_DOCX
    LANG, RTL, T = lang, lang == "he", STR[lang]
    FONT = T["font"]
    figures.configure(lang)
    if lang == "he":
        import content_he
        C = content_he
        OUT_DOCX = os.path.join(HERE, "API_Gateway_Enterprise_Architecture_Standard_he.docx")


# ---------------------------------------------------------------------------
# Low-level XML helpers
# ---------------------------------------------------------------------------

def insert_ordered(parent, child, order):
    """Insert child into parent at the schema-correct position."""
    tag = child.tag.split("}")[1]
    for existing in parent.findall(child.tag):        # replace same-tag element
        parent.remove(existing)
    idx = order.index(tag)
    for el in list(parent):
        name = el.tag.split("}")[1]
        if name in order and order.index(name) > idx:
            el.addprevious(child)
            return child
    parent.append(child)
    return child


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def para_border(pPr, side, color, sz, space=1):
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = insert_ordered(pPr, el("w:pBdr"), PPR_ORDER)
    for old in bdr.findall(qn(f"w:{side}")):
        bdr.remove(old)
    order = ["top", "left", "bottom", "right", "between", "bar"]
    new = el(f"w:{side}", val="single", sz=sz, space=space, color=color)
    for sib in list(bdr):
        if order.index(sib.tag.split("}")[1]) > order.index(side):
            sib.addprevious(new)
            return
    bdr.append(new)


def para_shade(pPr, fill):
    insert_ordered(pPr, el("w:shd", val="clear", color="auto", fill=fill), PPR_ORDER)


def set_style_font(style, name=None, size=None, bold=None, italic=None, color=None):
    name = name or FONT
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        if qn(f"w:{a}") in rfonts.attrib:
            del rfonts.attrib[qn(f"w:{a}")]
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{a}"), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def get_or_add_style(doc, name, base="Normal"):
    styles = doc.styles
    try:
        return styles[name]
    except KeyError:
        s = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if base:
            s.base_style = styles[base]
        return s


# ---------------------------------------------------------------------------
# Numbering pass and reference resolution
# ---------------------------------------------------------------------------

class Nums:
    ch, sec, fig, tbl = {}, {}, {}, {}
    toc, figlist, tbllist, changes = [], [], [], []


def number_document():
    n = 0
    tbl_n = fig_n = 0
    for ch in C.CHAPTERS:
        if ch["appendix"]:
            label = f"{T['appendix']} {T['appendix_letters'].get(ch['appendix'], ch['appendix'])}"
            prefix = ch["appendix"]
            ch["heading"] = f"{label} — {ch['title']}"
        else:
            n += 1
            label = f"{T['chapter']} {n}"
            prefix = str(n)
            ch["heading"] = f"{n}\t{ch['title']}"
        ch["label"], ch["prefix"] = label, prefix
        Nums.ch[ch["code"]] = label
        Nums.toc.append((1, ch["heading"]))
        sec_i, cur_sec = 0, None
        for b in ch["blocks"]:
            t = b["t"]
            if b.get("chg"):
                Nums.changes.append((cur_sec or f"{ch['label']} {ch['title']}", b["chg"]))
            if t == "h2":
                sec_i += 1
                b["num"] = f"{prefix}.{sec_i}"
                cur_sec = f"{b['num']} {b['title']}"
                Nums.sec[b["key"]] = f"{T['section']} {b['num']}"
                Nums.toc.append((2, f"{b['num']}\t{b['title']}"))
            elif t == "tbl":
                tbl_n += 1
                b["no"] = tbl_n
                Nums.tbl[b["key"]] = f"{T['table']} {tbl_n}"
                Nums.tbllist.append(f"{T['table']} {tbl_n} — {b['caption']}")
            elif t == "fig":
                fig_n += 1
                b["no"] = fig_n
                Nums.fig[b["key"]] = f"{T['figure']} {fig_n}"
                Nums.figlist.append(f"{T['figure']} {fig_n} — {b['caption']}")


REF = re.compile(r"\{(ch|sec|fig|tbl):(\w+)\}")


def resolve(s):
    def rep(m):
        kind, key = m.groups()
        table = {"ch": Nums.ch, "sec": Nums.sec, "fig": Nums.fig, "tbl": Nums.tbl}[kind]
        if key not in table:
            errors.append(f"unresolved reference {m.group(0)}")
            return m.group(0)
        return table[key]
    return REF.sub(rep, s).replace(f", {T['section']} ", ", ")


# ---------------------------------------------------------------------------
# Inline text, fields, lists, tables
# ---------------------------------------------------------------------------

TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|\[\[.+?\]\])")
PREFIX_HYPHEN = re.compile(r"(?<=[\u05D0-\u05EA])-(?=[A-Za-z0-9])")      # Hebrew prefix + hyphen + Latin word


def add_rtl_text(run, text):
    """Text of a Hebrew run. A prefix joined to a Latin word (ה-Gateway) gets a non-breaking hyphen, so
    that a line break cannot leave the prefix stranded at the end of a line."""
    for i, seg in enumerate(PREFIX_HYPHEN.split(text)):
        if i:
            run._r.append(OxmlElement("w:noBreakHyphen"))
        if seg:
            run.add_text(seg)


def add_inline(par, text, size=None, bold=None, color=None):
    text = resolve(text)
    for tok in TOKEN.split(text):
        if not tok:
            continue
        run_kw = {}
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            tok, run_kw["bold"] = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, run_kw["italic"] = tok[1:-1], True
        elif tok.startswith("`") and tok.endswith("`"):
            tok, run_kw["mono"] = tok[1:-1], True
        elif tok.startswith("[[") and tok.endswith("]]"):
            tok, run_kw["hl"] = "[" + tok[2:-2] + "]", True
        if RTL and "\t" not in tok and "\n" not in tok:
            run = par.add_run()
            add_rtl_text(run, tok)
        else:
            run = par.add_run(tok)
        if size:
            run.font.size = Pt(size)
        if bold or run_kw.get("bold"):
            run.font.bold = True
        if run_kw.get("italic"):
            run.font.italic = True
        if color is not None:
            run.font.color.rgb = color
        if run_kw.get("mono"):
            run.font.name = MONO
            run.font.size = Pt((size or 10.5) - 1)
        if run_kw.get("hl"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def add_field(par, instr, cached, size=None, bold=None):
    def marker(kind):
        r = par.add_run()
        r._r.append(el("w:fldChar", fldCharType=kind))
    marker("begin")
    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    marker("separate")
    run = par.add_run(cached)
    if size:
        run.font.size = Pt(size)
    if bold:
        run.font.bold = True
    marker("end")


def add_field_block(doc, instr, entries):
    """A multi-paragraph field (TOC / table of figures) with cached entries.
    entries: [(style_name, text)]"""
    paras = []
    for style, text in entries:
        p = doc.add_paragraph(style=style)
        paras.append(p)
        for i, part in enumerate(text.split("\t")):
            if i:
                p.add_run("\t")
            p.add_run(part)
    if not paras:
        return
    first, last = paras[0], paras[-1]
    begin = OxmlElement("w:r")
    begin.append(el("w:fldChar", fldCharType="begin"))
    ins = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    ins.append(it)
    sep = OxmlElement("w:r")
    sep.append(el("w:fldChar", fldCharType="separate"))
    anchor = first._p.find(qn("w:pPr"))
    for r in (sep, ins, begin):                         # insert after pPr, in order begin, instr, separate
        anchor.addnext(r)
    end = OxmlElement("w:r")
    end.append(el("w:fldChar", fldCharType="end"))
    last._p.append(end)


def add_bullet(container, text, left_cm, hang_cm=0.5, size=None, keep_next=False, style="List Bullet",
               space_after=2):
    p = container.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent = Cm(left_cm), Cm(-hang_cm)
    pf.space_after = Pt(space_after)
    pf.keep_with_next = keep_next
    add_inline(p, text, size=size)
    return p


def set_cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    order = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
             "textDirection", "tcFitText", "vAlign", "hideMark"]
    insert_ordered(tcPr, el("w:shd", val="clear", color="auto", fill=fill), order)


def fill_cell(cell, value, header=False, bold=False):
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)   # remove default empty paragraph
    if isinstance(value, list):
        for item in value:
            add_bullet(cell, item, 0.42, 0.34, size=9.5, style="List Bullet", space_after=1)
        return
    lines = str(value).split("\n") if value != "" else [""]
    for ln in lines:
        p = cell.add_paragraph(style="Table Text")
        add_inline(p, ln, bold=header or bold, color=RGBColor(0xFF, 0xFF, 0xFF) if header else None)


def add_table(doc, headers, rows, widths, first_bold=True):
    ncols = len(widths)
    body_rows = len(rows)
    tbl = doc.add_table(rows=body_rows + (1 if headers else 0), cols=ncols)
    if not RTL:                                        # RTL tables start at the right margin by default
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    if RTL:
        insert_ordered(tblPr, el("w:bidiVisual"), TBLPR_ORDER)   # first column on the right
    total = int(sum(widths) * 567)
    insert_ordered(tblPr, el("w:tblW", w=total, type="dxa"), TBLPR_ORDER)
    borders = el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(el(f"w:{side}", val="single", sz=4, space=0, color=GRID_HEX))
    insert_ordered(tblPr, borders, TBLPR_ORDER)
    insert_ordered(tblPr, el("w:tblLayout", type="fixed"), TBLPR_ORDER)
    mar = el("w:tblCellMar")
    for side, w in (("top", 50), ("left", 100), ("bottom", 50), ("right", 100)):
        mar.append(el(f"w:{side}", w=w, type="dxa"))
    insert_ordered(tblPr, mar, TBLPR_ORDER)
    for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 567)))

    all_rows = ([headers] if headers else []) + rows
    for ri, (row, data) in enumerate(zip(tbl.rows, all_rows)):
        is_head = bool(headers) and ri == 0
        trPr = row._tr.get_or_add_trPr()
        trPr.append(el("w:cantSplit"))
        if is_head:
            trPr.append(el("w:tblHeader"))
        for ci, (cell, val) in enumerate(zip(row.cells, data)):
            cell.width = Cm(widths[ci])
            fill_cell(cell, val, header=is_head, bold=(ci == 0 and first_bold and not is_head))
            if is_head:
                if RTL:
                    for par in cell.paragraphs:
                        par.paragraph_format.keep_with_next = True
                set_cell_shade(cell, NAVY_HEX)
            elif ci == 0 and first_bold:
                set_cell_shade(cell, "F2F5FA")
    return tbl


def caption(doc, kind, no, text, keep_next):
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_with_next = keep_next
    p.add_run(f"{T[kind.lower()]} ")                  # visible label; the SEQ identifier stays "Figure"/"Table"
    add_field(p, f"SEQ {kind} \\* ARABIC", str(no))
    add_inline(p, f" — {text}")
    return p


# ---------------------------------------------------------------------------
# Styles, page setup, header and footer
# ---------------------------------------------------------------------------

def setup_styles(doc):
    st = doc.styles
    normal = st["Normal"]
    set_style_font(normal, FONT, 10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.widow_control = True

    h1 = st["Heading 1"]
    set_style_font(h1, FONT, 17, bold=True, color=NAVY)
    h1.paragraph_format.space_before, h1.paragraph_format.space_after = Pt(26), Pt(10)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.left_indent, h1.paragraph_format.first_line_indent = Cm(1.3), Cm(-1.3)
    h1.paragraph_format.tab_stops.add_tab_stop(Cm(1.3))
    para_border(h1.element.get_or_add_pPr(), "bottom", BLUE_HEX, 8, space=3)

    h2 = st["Heading 2"]
    set_style_font(h2, FONT, 13, bold=True, color=BLUE)
    h2.paragraph_format.space_before, h2.paragraph_format.space_after = Pt(16), Pt(6)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.left_indent, h2.paragraph_format.first_line_indent = Cm(1.3), Cm(-1.3)
    h2.paragraph_format.tab_stops.add_tab_stop(Cm(1.3))

    cap = st["Caption"]
    set_style_font(cap, FONT, 9.5, bold=True, italic=False, color=NAVY)
    cap.paragraph_format.space_before, cap.paragraph_format.space_after = Pt(8), Pt(4)

    lb = st["List Bullet"]
    set_style_font(lb, FONT, 10.5)
    ppr = lb.element.get_or_add_pPr()
    for cs in ppr.findall(qn("w:contextualSpacing")):    # template suppresses spacing between bullets
        ppr.remove(cs)

    note = get_or_add_style(doc, "Note")
    note.paragraph_format.space_before, note.paragraph_format.space_after = Pt(4), Pt(8)
    note.paragraph_format.left_indent = Cm(0.3)
    note.paragraph_format.keep_together = True
    pPr = note.element.get_or_add_pPr()
    para_border(pPr, "right" if RTL else "left", BLUE_HEX, 24, space=6)   # bar on the start side
    para_shade(pPr, LIGHT_HEX)

    tt = get_or_add_style(doc, "Table Text")
    set_style_font(tt, FONT, 9.5)
    tt.paragraph_format.space_after, tt.paragraph_format.line_spacing = Pt(2), 1.05

    tocw = get_or_add_style(doc, "Front Heading")
    set_style_font(tocw, FONT, 17, bold=True, color=NAVY)
    tocw.paragraph_format.space_before, tocw.paragraph_format.space_after = Pt(4), Pt(10)
    tocw.paragraph_format.keep_with_next = True
    para_border(tocw.element.get_or_add_pPr(), "bottom", BLUE_HEX, 8, space=3)

    sub = get_or_add_style(doc, "Front Subheading")
    set_style_font(sub, FONT, 12, bold=True, color=BLUE)
    sub.paragraph_format.space_before, sub.paragraph_format.space_after = Pt(14), Pt(4)
    sub.paragraph_format.keep_with_next = True

    # Word's built-in TOC styles are looked up by these lower-case names.
    for name, left, tab in (("toc 1", 0, 1.3), ("toc 2", 0.6, 1.9), ("table of figures", 0, None)):
        s = get_or_add_style(doc, name)
        set_style_font(s, FONT, 10.5 if name != "toc 2" else 10, bold=(name == "toc 1"))
        s.paragraph_format.left_indent = Cm(left)
        s.paragraph_format.space_after = Pt(3 if name != "toc 1" else 4)
        s.paragraph_format.space_before = Pt(6 if name == "toc 1" else 0)
        if tab:
            s.paragraph_format.tab_stops.add_tab_stop(Cm(tab))
        s.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    for name in ("Header", "Footer"):
        s = st[name]
        set_style_font(s, FONT, 9, color=GRAY)
        s.paragraph_format.tab_stops.clear_all()
        s.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    # document language
    rpr = normal.element.get_or_add_rPr()
    if RTL:
        rpr.append(el("w:lang", val="he-IL", eastAsia="he-IL", bidi="he-IL"))
        insert_ordered(normal.element.get_or_add_pPr(), el("w:bidi"), PPR_ORDER)   # all styles inherit it
    else:
        rpr.append(el("w:lang", val="en-GB", eastAsia="en-GB"))


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.3)
    sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.1)


def tabbed_paragraph(par, left):
    add_inline(par, left)
    par.add_run("\t")


def fill_header_footer(sec0, sec1):
    hp = sec0.header.paragraphs[0]
    tabbed_paragraph(hp, C.META["title"])
    add_inline(hp, C.META["classification"])
    para_border(hp._p.get_or_add_pPr(), "bottom", GRID_HEX, 4, space=4)

    for sec, total in ((sec0, False), (sec1, True)):
        if sec is sec1:
            sec.footer.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]
        tabbed_paragraph(fp, T["footer_version"].format(v=C.META["version"], status=C.META["status"]))
        fp.add_run(T["page"] + " ")
        add_field(fp, "PAGE", "1")
        if total:
            fp.add_run(T["of"])
            add_field(fp, "SECTIONPAGES", "1")
        para_border(fp._p.get_or_add_pPr(), "top", GRID_HEX, 4, space=4)


def set_section_numbering(sec, fmt, start=None):
    sectPr = sec._sectPr
    attrs = {"fmt": fmt}
    if start is not None:
        attrs["start"] = start
    insert_ordered(sectPr, el("w:pgNumType", **attrs), SECTPR_ORDER)


# ---------------------------------------------------------------------------
# Front matter
# ---------------------------------------------------------------------------

def spacer(doc, pts):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)
    return p


def build_cover(doc):
    meta = C.META
    spacer(doc, 60)
    p = doc.add_paragraph()
    add_inline(p, meta["org"], size=14, bold=True, color=GRAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(30), Pt(6)
    add_inline(p, meta["title"], size=32, bold=True, color=NAVY)
    para_border(p._p.get_or_add_pPr(), "bottom", BLUE_HEX, 12, space=10)
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(80)
    add_inline(p, meta["subtitle"], size=16, color=GRAY)
    add_table(doc, None,
              [[T["version"], f"{meta['version']}"], [T["status"], meta["status"]], [T["date"], meta["date"]],
               [T["classification"], meta["classification"]]],
              [4.0, 8.0])
    doc.add_page_break()


def build_document_control(doc):
    meta = C.META
    doc.add_paragraph(T["doc_control"], style="Front Heading")
    add_table(doc, None,
              [[T["doc_title"], meta["title"]],
               [T["version"], meta["version"]], [T["status"], meta["status"]],
               [T["classification"], meta["classification"]], [T["owner"], meta["owner"]],
               [T["author"], meta["author"]], [T["effective"], meta["effective"]],
               [T["review_cycle"], meta["review_cycle"]], [T["next_review"], meta["next_review"]]],
              [4.6, 11.4])
    doc.add_paragraph(T["revision_history"], style="Front Subheading")
    add_table(doc, T["rev_headers"],
              [[r["version"], r["date"], r["author"], r["desc"]] for r in reversed(C.REVISIONS)],
              [2.2, 2.8, 3.6, 7.4], first_bold=False)
    doc.add_paragraph(T["approval"], style="Front Subheading")
    add_table(doc, T["appr_headers"],
              [[T["role"], T["name"], "", ""], [T["role"], T["name"], "", ""],
               [T["role"], T["name"], "", ""]],
              [4.6, 4.6, 4.0, 2.8], first_bold=False)
    doc.add_page_break()


def build_contents(doc):
    doc.add_paragraph(T["contents"], style="Front Heading")
    add_field_block(doc, 'TOC \\o "1-2" \\h \\z \\u',
                    [(f"toc {lvl}", text) for lvl, text in Nums.toc])
    doc.add_paragraph(T["list_figures"], style="Front Subheading").paragraph_format.page_break_before = True
    add_field_block(doc, 'TOC \\h \\z \\c "Figure"', [("table of figures", t) for t in Nums.figlist])
    doc.add_paragraph(T["list_tables"], style="Front Subheading")
    add_field_block(doc, 'TOC \\h \\z \\c "Table"', [("table of figures", t) for t in Nums.tbllist])


# ---------------------------------------------------------------------------
# Body
# ---------------------------------------------------------------------------

def render_picture(doc, b):
    path = os.path.join(figures.OUT_DIR, b["image"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(path, width=Cm(b["width_cm"]))
    docpr = run._r.xpath(".//wp:docPr")[0]
    docpr.set("descr", b["alt"])
    docpr.set("title", b["caption"])
    cp = caption(doc, "Figure", b["no"], b["caption"], keep_next=False)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(10)


def build_body(doc):
    for ch in C.CHAPTERS:
        h = doc.add_paragraph(style="Heading 1")
        if ch["appendix"]:
            h.paragraph_format.page_break_before = True
            h.paragraph_format.left_indent = h.paragraph_format.first_line_indent = Cm(0)
        h.paragraph_format.space_before = Pt(0 if (ch["prefix"] == "1" or ch["appendix"]) else 26)
        h.add_run(ch["heading"])
        for b in ch["blocks"]:
            t = b["t"]
            if t == "h2":
                hp = doc.add_paragraph(style="Heading 2")
                if ch["appendix"]:
                    hp.paragraph_format.left_indent = Cm(1.3)
                hp.add_run(f"{b['num']}\t{b['title']}")
            elif t == "p":
                add_inline(doc.add_paragraph(), b["text"])
            elif t == "note":
                p = doc.add_paragraph(style="Note")
                if not b["text"].startswith("**"):
                    p.add_run(T["note"]).font.bold = True
                add_inline(p, b["text"])
            elif t == "bul":
                for item in b["items"]:
                    add_bullet(doc, item, 0.9, 0.5, space_after=3)
            elif t == "tbl":
                caption(doc, "Table", b["no"], b["caption"], keep_next=True)
                add_table(doc, b["headers"], b["rows"], b["widths"], first_bold=b["first_bold"])
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            elif t == "fig":
                render_picture(doc, b)


# ---------------------------------------------------------------------------
# Document-level settings
# ---------------------------------------------------------------------------

def enable_update_fields(doc):
    settings = doc.settings.element
    zoom = settings.find(qn("w:zoom"))                  # template ships zoom without the required percent
    if zoom is not None:
        for a in list(zoom.attrib):
            del zoom.attrib[a]
        zoom.set(qn("w:percent"), "100")
    node = el("w:updateFields", val="true")
    compat = settings.find(qn("w:compat"))
    (compat.addprevious(node) if compat is not None else settings.append(node))


def set_properties(doc):
    cp = doc.core_properties
    cp.title, cp.subject = C.META["title"], C.META["subtitle"]
    cp.author = cp.last_modified_by = ""
    cp.keywords = T["keywords"]
    cp.created = cp.modified = datetime.fromisoformat(C.META["date_iso"]).replace(hour=12)
    cp.comments = T["comments"]
    if RTL:
        cp.language = "he-IL"


def strip_template_thumbnail(path):
    """The python-docx template carries a thumbnail of a blank page; drop it so the
    file preview is not misleading (Word regenerates it on save)."""
    import zipfile
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/thumbnail.jpeg":
                continue
            if item.filename == "_rels/.rels":
                data = re.sub(rb"<Relationship [^>]*thumbnail[^>]*/>", b"", data)
            zout.writestr(item, data)
    os.replace(tmp, path)


# ---------------------------------------------------------------------------
# Right-to-left finishing (Hebrew edition)
# ---------------------------------------------------------------------------

RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
             "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
             "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
             "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]


def normalize_rpr(rpr):
    """Give a run-properties element its complex-script twins and put its children in schema order.
    Word formats Hebrew with the complex-script properties (bCs, iCs, szCs), not with b, i and sz."""
    for base, twin in (("b", "bCs"), ("i", "iCs"), ("sz", "szCs")):
        src = rpr.find(qn(f"w:{base}"))
        if src is not None and rpr.find(qn(f"w:{twin}")) is None:
            new = OxmlElement(f"w:{twin}")
            for k, v in src.attrib.items():
                new.set(k, v)
            rpr.append(new)

    def rank(e):
        name = e.tag.split("}")[1]
        return RPR_ORDER.index(name) if name in RPR_ORDER else len(RPR_ORDER)
    for child in sorted(rpr, key=rank):
        rpr.append(child)                              # re-appending moves the element


def rtlify(doc):
    """Mark Hebrew runs as right-to-left, add complex-script formatting, set each section's direction."""
    roots = [doc.element, doc.styles.element]
    for sec in doc.sections:
        insert_ordered(sec._sectPr, el("w:bidi"), SECTPR_ORDER)
        for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
            if not part.is_linked_to_previous:
                roots.append(part._element)
    for root in roots:
        for r in root.iter(qn("w:r")):
            if HEBREW.search("".join(t.text or "" for t in r.iter(qn("w:t")))):
                rpr = r.get_or_add_rPr()
                if rpr.find(qn("w:rtl")) is None:
                    rpr.append(OxmlElement("w:rtl"))
        for rpr in root.iter(qn("w:rPr")):
            normalize_rpr(rpr)


def check_parity():
    """A translation must have exactly the structure of the English content: same chapters, blocks,
    keys, source sections, table shapes, figures, cross-references and placeholders."""
    E = C_EN

    def texts(b):
        t = b["t"]
        if t in ("p", "note"):
            return [b["text"]]
        if t == "bul":
            return list(b["items"])
        if t == "h2":
            return [b["title"]]
        if t == "fig":
            return [b["caption"], b["alt"]]
        cells = list(b["headers"] or [])
        for row in b["rows"]:
            for c in row:
                cells.extend(c if isinstance(c, list) else [c])
        return [b["caption"]] + cells

    def marks(text):
        return sorted(REF.findall(text)), text.count("[[")

    if set(C.META) != set(E.META):
        errors.append("parity: META keys differ")
    if [r["version"] for r in C.REVISIONS] != [r["version"] for r in E.REVISIONS]:
        errors.append("parity: REVISIONS versions differ")
    if [c["code"] for c in C.CHAPTERS] != [c["code"] for c in E.CHAPTERS]:
        errors.append("parity: chapter codes differ")
        return
    for ce, ct in zip(E.CHAPTERS, C.CHAPTERS):
        code = ce["code"]
        if (ce["appendix"], ce["src"]) != (ct["appendix"], ct["src"]):
            errors.append(f"parity: {code}: appendix letter or README sources differ")
        if [b["t"] for b in ce["blocks"]] != [b["t"] for b in ct["blocks"]]:
            errors.append(f"parity: {code}: block sequence differs")
            continue
        for i, (be, bt) in enumerate(zip(ce["blocks"], ct["blocks"]), 1):
            spot, t = f"{code} block {i}", be["t"]
            fields = {"h2": ("key", "src"), "tbl": ("key", "widths", "first_bold"),
                      "fig": ("key", "image", "width_cm")}.get(t, ())
            for f in fields:
                if be[f] != bt[f]:
                    errors.append(f"parity: {spot}: '{f}' differs")
            if t == "tbl" and ([len(r) for r in be["rows"]] != [len(r) for r in bt["rows"]]
                               or len(be["headers"] or []) != len(bt["headers"] or [])
                               or [[isinstance(c, list) and len(c) for c in r] for r in be["rows"]]
                               != [[isinstance(c, list) and len(c) for c in r] for r in bt["rows"]]):
                errors.append(f"parity: {spot}: table shape differs")
            te, tt = texts(be), texts(bt)
            if len(te) != len(tt):
                errors.append(f"parity: {spot}: {len(tt)} text items, expected {len(te)}")
                continue
            for a, b in zip(te, tt):
                if marks(a) != marks(b):
                    errors.append(f"parity: {spot}: cross-references or placeholders differ: {b[:50]!r}")


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------

def build():
    if RTL:
        check_parity()
        if errors:
            sys.exit("BUILD ERRORS:\n  - " + "\n  - ".join(errors))
    figures.generate_all()
    number_document()
    doc = Document()
    setup_page(doc)
    setup_styles(doc)

    build_cover(doc)
    build_document_control(doc)
    build_contents(doc)
    doc.add_section(WD_SECTION.NEW_PAGE)
    build_body(doc)

    sec0, sec1 = doc.sections[0], doc.sections[1]
    sec0.different_first_page_header_footer = True     # clean cover page
    sec1.different_first_page_header_footer = False
    set_section_numbering(sec0, "lowerRoman", start=1)
    set_section_numbering(sec1, "decimal", start=1)
    fill_header_footer(sec0, sec1)
    enable_update_fields(doc)
    set_properties(doc)
    if RTL:
        rtlify(doc)

    if errors:
        print("BUILD ERRORS:")
        for e in errors:
            print("  -", e)
        sys.exit(1)
    doc.save(OUT_DOCX)
    strip_template_thumbnail(OUT_DOCX)
    return doc


# ---------------------------------------------------------------------------
# CHANGES.md
# ---------------------------------------------------------------------------

def readme_sections():
    out = {}
    with open(README, encoding="utf-8") as f:
        for line in f:
            m = re.match(r"^#{1,2}\s+(\d+)\.\s+(.+?)\s*$", line)
            if m:
                out[int(m.group(1))] = m.group(2)
    return out


def write_changes():
    titles = readme_sections()
    where = {}
    for ch in C.CHAPTERS:
        for s in ch["src"]:
            where.setdefault(s, []).append(f"{ch['label']} {ch['title']}")
        for b in ch["blocks"]:
            if b["t"] == "h2":
                for s in b["src"]:
                    where.setdefault(s, []).append(f"{b['num']} {b['title']}")
    missing = [i for i in range(1, 41) if i not in where]
    if missing:
        errors.append(f"README sections not mapped: {missing}")

    L = [f"# Change log: README.md → API Gateway Enterprise Architecture Standard (version {C.META['version']})", "",
         "Generated by `build_docx.py`. Review this file first: it lists every place where the rewrite "
         "changed meaning, not only wording.", "",
         "## How the wording was converted", "",
         "From version 1.2 the document is continuous prose, without numbered requirements. Obligations "
         "are carried by the wording of the sentence:", "",
         "- **must / must not** ← README *must*, *required*, *mandatory*, *not permitted*, and items marked "
         "'Required' or 'separate' in the README capability summary (§38).",
         "- **should / is recommended** ← README *should* where it expresses a preference or recommendation.",
         "- **can / may** ← README *can*, *may*, *optional*.",
         "- **Present tense** ← descriptive statements in the README (for example 'The application remains "
         "responsible for…', 'The Gateway validates…'), which set out how the platform is required to operate.", "",
         "Versions 1.0 and 1.1 (in `versions/`) use numbered requirements with *shall*, *should* and *may*.", "",
         "## Places where the meaning differs from the README", "",
         "Almost all are a README 'should' strengthened to 'must' or to the present tense. Please confirm "
         "each, or downgrade it to 'should'. The security-core rules (C2B/B2B tokens, no reliance on "
         "undocumented manual production configuration) were made mandatory by a review decision; other "
         "'should' statements are unchanged.", "",
         "| Section | Change |", "|---|---|"]
    for section, chg in Nums.changes:
        L.append(f"| {section} | {resolve(chg)} |")
    L += ["", "## Inconsistencies in the README and how they were resolved", ""]
    L += [f"{i}. {resolve(t)}" for i, t in enumerate(C.RESOLVED, 1)]
    L += ["", "## New content (not in the README)", ""] + [f"- {resolve(t)}" for t in C.NEW_CONTENT]
    L += ["", "## Decisions taken during review", "",
          "Open issues from the first draft that were settled. Each is now written into the document.", "",
          "| Topic | Decision and where it is recorded |", "|---|---|"]
    L += [f"| {topic} | {resolve(text)} |" for topic, text in C.DECISIONS]
    L += ["", "## Where each README section went", "",
          "| README § | README title | Location in the standard |", "|---|---|---|"]
    appx = {}
    for i in range(1, 41):
        loc = "; ".join(where.get(i, [])) or appx.get(i, "—")
        L.append(f"| {i} | {titles.get(i, '')} | {loc} |")
    L += ["", "## Placeholders to complete", "",
          "Highlighted yellow in the document: organization name, document ID, classification, document "
          "owner, author, effective date, review cycle, next review date, approver roles and names, and "
          "the two organizational references in Section 1.6.", "",
          "## Decisions the README does not make", "",
          f"See {resolve('{ch:APC}')} of the document (OI-01).", ""]
    with open(OUT_CHANGES, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


# ---------------------------------------------------------------------------
# Version archive
# ---------------------------------------------------------------------------

VERSIONS_DIR = os.path.join(HERE, "versions")
DOC_STEM = "API_Gateway_Enterprise_Architecture_Standard"


def write_index():
    """versions/INDEX.md: one row per archived version, from C.REVISIONS."""
    rows = []
    for r in reversed(C.REVISIONS):
        folder = f"v{r['version']}"
        if os.path.isdir(os.path.join(VERSIONS_DIR, folder)):
            he = f"{folder}/{DOC_STEM}_he_{folder}.docx"
            he_link = f" ([עברית]({he}))" if os.path.exists(os.path.join(VERSIONS_DIR, he)) else ""
            rows.append(f"| {r['version']} | {r['date']} | [{folder}/{DOC_STEM}_{folder}.docx]"
                        f"({folder}/{DOC_STEM}_{folder}.docx){he_link} | "
                        f"{resolve(r['desc']).replace('[', '').replace(']', '')} |")
    text = ["# Document versions", "",
            "Every issued version of the API Gateway Enterprise Architecture Standard. Archived versions "
            "are never overwritten; a change is issued as a new version.", "",
            "| Version | Date | File | Summary |", "|---|---|---|---|", *rows, "",
            "Each folder holds the .docx and the `CHANGES.md` that was current when it was issued. "
            "The latest working copy is `../" + DOC_STEM + ".docx`.", ""]
    with open(os.path.join(VERSIONS_DIR, "INDEX.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(text))


def archive():
    ver = C.META["version"]
    if not C.REVISIONS or C.REVISIONS[-1]["version"] != ver:
        print(f"cannot archive: latest REVISIONS entry is not version {ver}")
        return 1
    dest = os.path.join(VERSIONS_DIR, f"v{ver}")
    if RTL:
        # The Hebrew edition joins the folder created by the English archive (run that first).
        target = os.path.join(dest, f"{DOC_STEM}_he_v{ver}.docx")
        if not os.path.isdir(dest):
            print(f"cannot archive: {dest} does not exist; archive the English edition first.")
            return 1
        if os.path.exists(target):
            print(f"cannot archive: {target} already exists. Archived versions are not overwritten.")
            return 1
        shutil.copy(OUT_DOCX, target)
        write_index()
        print(f"archived Hebrew edition {ver} -> {target}")
        return 0
    if os.path.exists(dest):
        print(f"cannot archive: {dest} already exists. Archived versions are not overwritten; "
              "bump META['version'] and add a REVISIONS entry.")
        return 1
    os.makedirs(dest)
    shutil.copy(OUT_DOCX, os.path.join(dest, f"{DOC_STEM}_v{ver}.docx"))
    shutil.copy(OUT_CHANGES, os.path.join(dest, "CHANGES.md"))
    write_index()
    print(f"archived version {ver} -> {dest}")
    return 0


# ---------------------------------------------------------------------------
# Verification of the saved file
# ---------------------------------------------------------------------------

def verify():
    doc = Document(OUT_DOCX)
    problems = []
    styles = [p.style.name for p in doc.paragraphs]

    # heading order: no H2 without a preceding H1
    seen_h1 = False
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            seen_h1 = True
        elif p.style.name == "Heading 2" and not seen_h1:
            problems.append(f"H2 before any H1: {p.text}")

    # wording: prose document, no numbered requirements, no 'shall'
    body_text = " ".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                body_text += " " + cell.text
    if not RTL:
        if re.search(r"\bshall\b", body_text, re.I):
            problems.append("the word 'shall' is still used")
        if re.search(r"\bGW-[A-Z]+-\d+", body_text):
            problems.append("requirement identifiers are still present")
    words = len(body_text.split())
    modals = {w: len(re.findall(rf"\b{w}\b", body_text)) for w in ("must", "should", "may", "can")}

    # figures: alt text and caption; tables: caption count
    pics = doc.element.body.xpath(".//wp:docPr")
    if not pics or any(not d.get("descr") for d in pics):
        problems.append("figure without alt text")
    cap_fig = sum(1 for p in doc.paragraphs if p.style.name == "Caption" and p.text.startswith(T["figure"]))
    cap_tbl = sum(1 for p in doc.paragraphs if p.style.name == "Caption" and p.text.startswith(T["table"]))
    if cap_fig != len(pics):
        problems.append(f"{len(pics)} figures but {cap_fig} figure captions")
    if cap_tbl != len(Nums.tbllist):
        problems.append(f"{len(Nums.tbllist)} captioned tables expected, {cap_tbl} found")

    # placeholders must be highlighted; no leftover markup or unresolved refs
    def all_runs():
        for p in doc.paragraphs:
            yield from p.runs
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        yield from p.runs
    for r in all_runs():
        if "[" in r.text and r.font.highlight_color is None:
            problems.append(f"unhighlighted placeholder: {r.text[:50]}")
        if REF.search(r.text) or "**" in r.text or "[[" in r.text:
            problems.append(f"leftover markup: {r.text[:50]}")
        if re.search(r"\b(TBD|TODO|XXX)\b", r.text):
            problems.append(f"unfinished marker: {r.text[:50]}")

    # every README section 1-40 is mapped
    mapped = set()
    for ch in C.CHAPTERS:
        mapped.update(ch["src"])
        for b in ch["blocks"]:
            if b["t"] == "h2":
                mapped.update(b["src"])
    if not set(range(1, 41)) <= mapped:
        problems.append(f"README sections not covered: {sorted(set(range(1, 41)) - mapped)}")

    if RTL:
        # direction: base style, every table, every section; and no paragraph left untranslated
        if doc.styles["Normal"].element.pPr.find(qn("w:bidi")) is None:
            problems.append("Normal style is not right-to-left")
        if any(t._tbl.tblPr.find(qn("w:bidiVisual")) is None for t in doc.tables):
            problems.append("a table is not right-to-left")
        if any(sec._sectPr.find(qn("w:bidi")) is None for sec in doc.sections):
            problems.append("a section is not right-to-left")
        for p in doc.paragraphs:                       # headings may be a bare product name (GitOps)
            heading = p.style.name.startswith(("Heading", "toc", "table of figures"))
            citation = p.text.lstrip("\u200e").startswith("BCP 14")   # external standard, cited in English
            if p.text.strip() and not heading and not citation and not HEBREW.search(p.text):
                problems.append(f"paragraph without Hebrew text: {p.text[:60]!r}")
    stats = f"words={words}" if RTL else f"words={words} modal verbs={modals}"
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} figures={len(pics)} "
          f"sections={len(doc.sections)} {stats}")
    if not C.REVISIONS or C.REVISIONS[-1]["version"] != C.META["version"]:
        problems.append("META version does not match the latest REVISIONS entry")
    if problems:
        print("VERIFY PROBLEMS:")
        for p in problems:
            print("  -", p)
        return 1
    print("verify: OK")
    return 0


if __name__ == "__main__":
    args = sys.argv[1:]
    lang = args[args.index("--lang") + 1] if "--lang" in args else "en"
    if lang not in STR:
        sys.exit(f"unknown language {lang!r}; available: {', '.join(STR)}")
    select_language(lang)
    if "--out" in args:
        OUT_DOCX = os.path.abspath(args[args.index("--out") + 1])
    build()
    if not RTL:
        write_changes()
    if errors:
        print("ERRORS:", *errors, sep="\n  - ")
        sys.exit(1)
    print(f"wrote {OUT_DOCX}")
    if not RTL:
        print(f"wrote {OUT_CHANGES}")
    status = verify()
    if status == 0 and "--archive" in sys.argv:
        status = archive()
    sys.exit(status)
